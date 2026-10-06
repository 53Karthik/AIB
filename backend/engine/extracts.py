import csv
import io
import math

from .outcome import norm


EXTRACT_KINDS = [
    {"kind": "ebq", "label": "EBQ Correspondence Report", "feeds": ["23B_NUL", "23B_UL_S1"],
     "headers": ["product", "policy number", "transaction reference", "transaction type", "status", "transaction start date", "transaction merged date"]},
    {"kind": "cancellation", "label": "Cancellation Tracker (CANREVEXT)", "feeds": ["23C"],
     "headers": ["policy number", "product name", "cancellation reason", "date and time of last status"]},
    {"kind": "withdrawal", "label": "Withdrawal extract (WITHDRAWALEXT)", "feeds": ["23B_UL_S2"],
     "headers": ["policy number", "product name", "request id", "transaction status", "transaction type", "date and time of last status"]},
    {"kind": "workflow", "label": "Workflow extract (WRKFLWEXT)", "feeds": ["23A", "23B_UL_S2", "23B_UL_S3", "23C", "23E"],
     "headers": ["workflow number", "workflow type", "reference number", "product name", "created date", "status", "close date", "workflow description"]},
]


def is_blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def text(record, field):
    value = record.get(field)
    if is_blank(value):
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def id_of(record, field):
    value = record.get(field)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(math.floor(value + 0.5))
    return text(record, field)


def _cell_value(value):
    if isinstance(value, bool):
        return str(value).lower()
    return value


def read_extract(buffer, filename):
    extension = str(filename).rsplit(".", 1)[-1].lower()
    if extension not in ("csv", "xlsx"):
        raise ValueError(f"Unsupported file type: .{extension} (expected .csv or .xlsx)")
    if extension == "csv":
        rows = csv.reader(io.StringIO(buffer.decode("utf-8-sig"), newline=""))
        grid = [(number, row) for number, row in enumerate(rows, 1)]
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(buffer), data_only=True, read_only=True)
        try:
            if {"Mapping for 23B", "Monthly summary"}.issubset(workbook.sheetnames):
                raise ValueError("This is the SLA expected-results workbook (rules, mapping and expected outcomes), not a BaNCS extract")
            grid = []
            for number, row in enumerate(workbook.worksheets[0].iter_rows(), 1):
                cells = [_cell_value(cell.value) if cell.data_type != "e" else None for cell in row]
                if any(value is not None for value in cells):
                    grid.append((number, cells))
        finally:
            workbook.close()
    found = None
    for number, cells in grid[:10]:
        names = set(map(norm, cells))
        kind = next((kind for kind in EXTRACT_KINDS if set(kind["headers"]) <= names), None)
        if kind:
            found = number, cells, kind
            break
    if found is None:
        raise ValueError("Columns do not match any BaNCS extract (EBQ, CANREVEXT, WITHDRAWALEXT or WRKFLWEXT)")
    header_row, headers, kind = found
    fields = list(map(norm, headers))
    records = []
    for number, cells in grid:
        if number <= header_row or all(is_blank(value) for value in cells):
            continue
        record = {"_row": number}
        for index, field in enumerate(fields):
            if field:
                record[field] = cells[index] if index < len(cells) else None
        records.append(record)
    return {
        "kind": kind["kind"], "label": kind["label"], "ext": extension,
        "title": [str(value).strip() for number, cells in grid if number < header_row for value in cells if not is_blank(value)],
        "headerRow": header_row,
        "columns": [str(value).strip() for value in headers if not is_blank(value)],
        "records": records,
    }
