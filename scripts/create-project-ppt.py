from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "AIB_Life_SLA_Governance_Project.pptx"
LOGO = ROOT / "public" / "aib-life-logo.jpg"

NAVY = RGBColor(24, 26, 57)
PURPLE = RGBColor(79, 57, 133)
LILAC = RGBColor(237, 232, 248)
CREAM = RGBColor(249, 248, 245)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(39, 37, 55)
MUTED = RGBColor(103, 98, 121)
GREEN = RGBColor(22, 132, 104)
RED = RGBColor(190, 61, 67)


def rect(slide, x, y, w, h, fill, radius=True, line=None):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def base(slide, number):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = CREAM
    rect(slide, 0, 0, 13.333, 0.12, PURPLE, radius=False)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(11.75), Inches(0.35), width=Inches(1.05))
    text(slide, f"AIB LIFE  /  SCHEDULE 23 GOVERNANCE", 0.55, 0.34, 7.2, 0.3, 9, MUTED, True)
    text(slide, f"0{number}", 12.25, 7.05, 0.5, 0.2, 9, MUTED, True, PP_ALIGN.RIGHT)


def title(slide, kicker, heading, subheading):
    text(slide, kicker.upper(), 0.65, 0.88, 3.5, 0.3, 11, PURPLE, True)
    text(slide, heading, 0.65, 1.25, 11.8, 0.75, 30, NAVY, True)
    text(slide, subheading, 0.67, 2.03, 11.4, 0.52, 14, MUTED)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# Slide 1 — Problem
slide = prs.slides.add_slide(blank)
base(slide, 1)
title(slide, "Problem statement", "Monthly SLA governance is slow, fragmented and difficult to evidence",
      "Schedule 23 performance must be reconstructed from multiple BaNCS extracts and complex business-day rules.")

problems = [
    ("01", "Fragmented inputs", "Five extract slots across EBQ, cancellations, withdrawals and open/closed workflows."),
    ("02", "Complex measurement", "Different targets, cut-offs, holidays, workflow matching and reporting-month rules."),
    ("03", "Weak auditability", "Manual summaries make it hard to trace a result back to the exact source record."),
]
for i, (n, head, body) in enumerate(problems):
    x = 0.65 + i * 4.13
    rect(slide, x, 2.85, 3.75, 2.25, WHITE, line=LILAC)
    text(slide, n, x + 0.28, 3.10, 0.5, 0.35, 12, PURPLE, True)
    text(slide, head, x + 0.28, 3.55, 3.1, 0.42, 18, NAVY, True)
    text(slide, body, x + 0.28, 4.08, 3.05, 0.75, 12, MUTED)

rect(slide, 0.65, 5.48, 12.0, 0.92, NAVY)
text(slide, "BUSINESS IMPACT", 0.95, 5.74, 1.6, 0.25, 10, RGBColor(190, 180, 223), True)
text(slide, "Delayed governance packs  •  inconsistent interpretation  •  limited confidence in reported figures",
     2.52, 5.68, 9.65, 0.36, 15, WHITE, True)

# Slide 2 — Solution
slide = prs.slides.add_slide(blank)
base(slide, 2)
title(slide, "Implemented solution", "One traceable pipeline from raw extracts to governance intelligence",
      "The application calculates every SLA result, builds monthly packs and explains performance across the full history.")

steps = [
    ("1", "INGEST", "CSV / XLSX\nextract set"),
    ("2", "CALCULATE", "Schedule 23\nrules engine"),
    ("3", "GOVERN", "Monthly packs\n& exceptions"),
    ("4", "UNDERSTAND", "Trends, drivers\n& guarded AI"),
]
for i, (n, head, body) in enumerate(steps):
    x = 0.65 + i * 3.05
    rect(slide, x, 2.72, 2.55, 1.55, WHITE, line=LILAC)
    rect(slide, x + 0.18, 2.92, 0.42, 0.42, PURPLE)
    text(slide, n, x + 0.18, 2.93, 0.42, 0.38, 12, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, head, x + 0.72, 2.92, 1.55, 0.26, 10, PURPLE, True)
    text(slide, body, x + 0.22, 3.43, 2.05, 0.57, 14, NAVY, True)
    if i < 3:
        text(slide, "→", x + 2.58, 3.22, 0.42, 0.35, 22, PURPLE, True, PP_ALIGN.CENTER)

features = [
    (GREEN, "Exact", "0-difference verification against the expected-results workbook"),
    (PURPLE, "Traceable", "Drill from every SLA result to its source item and deadline"),
    (RED, "Actionable", "Exceptions, backlog, failure drivers and pass/fail trends"),
]
for i, (accent, head, body) in enumerate(features):
    x = 0.65 + i * 4.13
    rect(slide, x, 4.72, 3.75, 1.25, LILAC, line=LILAC)
    rect(slide, x, 4.72, 0.09, 1.25, accent, radius=False)
    text(slide, head, x + 0.30, 4.94, 1.2, 0.3, 15, NAVY, True)
    text(slide, body, x + 0.30, 5.32, 3.08, 0.45, 11, MUTED)

text(slide, "React + Vite frontend  •  Node.js + Express API  •  deterministic rules engine  •  Amazon Nova Pro with guardrails",
     0.67, 6.36, 11.8, 0.35, 11, MUTED, False, PP_ALIGN.CENTER)

# Slide 3 — Video
slide = prs.slides.add_slide(blank)
base(slide, 3)
title(slide, "Product demonstration", "See the workflow end to end",
      "Use this slide for the recorded walkthrough of data ingestion, monthly governance and intelligence.")

video = rect(slide, 1.02, 2.70, 11.30, 3.55, NAVY)
video.line.color.rgb = PURPLE
video.line.width = Pt(1.5)
circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(5.91), Inches(3.63), Inches(1.5), Inches(1.5))
circle.fill.solid()
circle.fill.fore_color.rgb = WHITE
circle.line.color.rgb = WHITE
text(slide, "▶", 6.08, 3.88, 1.17, 0.66, 31, PURPLE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(slide, "INSERT PROJECT DEMO VIDEO", 4.10, 5.35, 5.15, 0.35, 14, WHITE, True, PP_ALIGN.CENTER)
text(slide, "Suggested flow: Extracts  →  Dashboard  →  SLA position  →  Exceptions  →  Intelligence",
     2.15, 6.57, 9.05, 0.32, 12, MUTED, False, PP_ALIGN.CENTER)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
